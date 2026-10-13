"""Exportación de reportes a Excel con OpenPyXL."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def build_report_xlsx(report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    header_fill = PatternFill("solid", fgColor="E85D26")
    header_font = Font(color="FFFFFF", bold=True)

    ws["A1"] = "SaborExpress - Operational Report"
    ws["A1"].font = Font(size=14, bold=True, color="E85D26")
    ws.merge_cells("A1:D1")

    ws["A2"] = f"From: {report['date_from']}  |  To: {report['date_to']}  |  Branch: {report['branch_id']}"
    ws.merge_cells("A2:D2")

    summary = [
        ("Total sales", report["total_sales"]),
        ("Total orders", report["total_orders"]),
        ("Average ticket", report["avg_ticket"]),
        ("IGV total (18%)", report["igv_total"]),
    ]
    row = 4
    for label, value in summary:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Breakdown by payment method").font = Font(bold=True)
    row += 1
    for col, title in enumerate(["Method", "Count", "Total"], start=1):
        c = ws.cell(row=row, column=col, value=title)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")
    row += 1
    for b in report["breakdown"]:
        ws.cell(row=row, column=1, value=b["method"])
        ws.cell(row=row, column=2, value=b["count"])
        ws.cell(row=row, column=3, value=b["total"])
        row += 1

    for col, width in zip("ABCD", (24, 16, 16, 16)):
        ws.column_dimensions[col].width = width

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
